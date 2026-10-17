"""Экспорт в PDF / DOCX из единого списка блоков — чтобы ученическая и учительская версии не расходились."""
import io
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

from .catalog import DIFFICULTIES, WORK_TYPES
from .services import student_view

FONTS = Path(__file__).parent / "static" / "fonts"
LETTERS = "АБВГДЕ"
_registered = False


def _fonts():
    global _registered
    if not _registered:
        pdfmetrics.registerFont(TTFont("DV", str(FONTS / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DV-B", str(FONTS / "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("DV", normal="DV", bold="DV-B", italic="DV", boldItalic="DV-B")
        _registered = True


def _num(x):
    return str(int(x)) if float(x).is_integer() else str(x).replace(".", ",")


def blocks(content, kind, created=""):
    """kind: 'student' | 'teacher'. Возвращает [(тип, текст)]."""
    m = content.get("meta", {})
    b = []
    if m.get("kind") == "lesson":
        return _lesson_blocks(content, created)
    src = student_view(content) if kind == "student" else content
    title = src["title"] + (" — ответы и решения (для учителя)" if kind == "teacher" else "")
    b.append(("h1", title))
    info = [f"Предмет: {m.get('subject', '')}", f"Класс: {m.get('grade', '')}{m.get('class_letter', '')}",
            f"Тема: {m.get('topic', '')}", f"Формат: {WORK_TYPES.get(m.get('kind'), '')}"]
    if m.get("duration_minutes"):
        info.append(f"Время: {m['duration_minutes']} мин.")
    if m.get("max_score"):
        info.append(f"Максимальный балл: {_num(m['max_score'])}")
    if m.get("difficulty") in DIFFICULTIES:
        info.append(f"Сложность: {DIFFICULTIES[m['difficulty']].lower()}")
    if created:
        info.append(f"Дата создания: {created}")
    b.append(("p", " · ".join(info)))
    if kind == "student":
        b.append(("p", "Фамилия, имя: ______________________________   Дата: ______________"))
    n = len(src["variants"])
    for vi, v in enumerate(content["variants"] if kind == "teacher" else src["variants"], 1):
        if vi > 1:
            b.append(("pb", ""))
        b.append(("h2", f"Вариант {vi}" if n > 1 or kind == "teacher" else "Задания"))
        for ti, t in enumerate(v["tasks"], 1):
            pts = f" ({_num(t['points'])} б.)" if t.get("points") else ""
            if kind == "student":
                b.append(("p", f"{ti}. {t['text']}{pts}"))
                if t.get("options"):
                    for oi, o in enumerate(t["options"]):
                        b.append(("li", f"{LETTERS[oi]}) {o}"))
                else:
                    b.append(("line", "Ответ: " + "_" * 60))
            else:
                if t.get("options"):
                    ans = f"{LETTERS[t['correct']]}) {t['options'][t['correct']]}" if t["options"] else ""
                else:
                    ans = t.get("answer", "")
                b.append(("p", f"{ti}. Ответ: {ans}" + ("  ⚠ проверьте вычисления" if t.get("verified") is False else "")))
                if t.get("solution"):
                    b.append(("li", "Решение: " + t["solution"]))
    if kind == "teacher" and content.get("criteria"):
        b.append(("pb", ""))
        b.append(("h2", "Критерии оценивания"))
        for c in content["criteria"]:
            b.append(("li", c))
    return b


def _lesson_blocks(c, created):
    m, b = c["meta"], []
    b += [("h1", c["title"]),
          ("p", f"Предмет: {m.get('subject', '')} · Класс: {m.get('grade', '')}{m.get('class_letter', '')} · "
                f"Тема: {m.get('topic', '')} · Продолжительность: {m.get('duration_minutes', '')} мин."
                + (f" · Дата создания: {created}" if created else ""))]
    b.append(("h2", "Цели урока"))
    b += [("li", g) for g in c["goals"]]
    b.append(("h2", "План урока"))
    b += [("li", f"{s['name']} — {s['minutes']} мин. {s['description']}") for s in c["stages"]]
    if c["explanation"]:
        b += [("h2", "Объяснение материала"), ("p", c["explanation"])]
    if c["examples"]:
        b.append(("h2", "Примеры"))
        b += [("li", f"{i}. {e}") for i, e in enumerate(c["examples"], 1)]
    if c["questions"]:
        b.append(("h2", "Вопросы для класса"))
        b += [("li", f"{i}. {q}") for i, q in enumerate(c["questions"], 1)]
    if c["practice"]:
        b.append(("h2", "Практические задания"))
        b += [("li", f"{i}. {p['text']}") for i, p in enumerate(c["practice"], 1)]
    if c["teacher_notes"] or any(p["answer"] for p in c["practice"]):
        b.append(("h2", "Ответы и пояснения для учителя"))
        b += [("li", f"{i}. {p['answer']}") for i, p in enumerate(c["practice"], 1) if p["answer"]]
        if c["teacher_notes"]:
            b.append(("p", c["teacher_notes"]))
    if c["homework"]:
        b += [("h2", "Домашнее задание"), ("p", c["homework"])]
    if c["summary"]:
        b += [("h2", "Краткий итог урока"), ("p", c["summary"])]
    return b


def to_pdf(blks) -> bytes:
    _fonts()
    st = {
        "h1": ParagraphStyle("h1", fontName="DV-B", fontSize=16, leading=21, spaceAfter=6, alignment=TA_LEFT),
        "h2": ParagraphStyle("h2", fontName="DV-B", fontSize=13, leading=17, spaceBefore=10, spaceAfter=6),
        "p": ParagraphStyle("p", fontName="DV", fontSize=10.5, leading=15, spaceAfter=5),
        "li": ParagraphStyle("li", fontName="DV", fontSize=10.5, leading=15, leftIndent=14, spaceAfter=3),
        "line": ParagraphStyle("line", fontName="DV", fontSize=10.5, leading=22, textColor="#555555", spaceAfter=8),
    }
    story = []
    for kind, text in blks:
        if kind == "pb":
            story.append(PageBreak())
        else:
            story.append(Paragraph(escape(text).replace("\n", "<br/>"), st[kind]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                      bottomMargin=16 * mm, title="TeacherAI").build(story or [Spacer(1, 1)])
    return buf.getvalue()


def to_docx(blks) -> bytes:
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    for kind, text in blks:
        if kind == "h1":
            d.add_heading(text, level=1)
        elif kind == "h2":
            d.add_heading(text, level=2)
        elif kind == "pb":
            d.add_page_break()
        elif kind == "li":
            d.add_paragraph(text).paragraph_format.left_indent = Pt(18)
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
